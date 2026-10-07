"""
common.py — 原始 Excel 读取与归一化

只做一件事：把 6 个 OTWS 导出文件读成干净的 Python 记录。
不做任何业务聚合，聚合在 datasets.py 里。

文件名说明：
  OTWS 导出的文件名没有业务含义（09.01-09.30.xlsx 这种），
  所以按 sheet 名识别文件类型，而不是按文件名：
    - 有「需求详情」+「发单详情」sheet  -> 需求池文件
    - 有「派遣人员名单」sheet            -> 派遣名单文件
    - 有「考勤」sheet                    -> 考勤文件
    - 有「临时工派遣」sheet              -> 临时工派遣文件
"""

import re
from pathlib import Path

import openpyxl

# 注意：不要用 read_only=True。OTWS 导出的文件用了 inlineStr 存单元格，
# openpyxl 3.1.2 的 read_only 流式解析读这种文件只吐出表头一行就停，
# 会误判成空文件（2026-10-06 实测）。普通模式读取正常。

# 收集数据问题，最后统一打印，不中断流程
warnings = []


def warn(msg):
    warnings.append(msg)


def reset_warnings():
    warnings.clear()


# ------------------------------------------------------------------
# 基础转换
# ------------------------------------------------------------------

def to_int(value):
    """1.0 -> 1, None/空 -> 0"""
    if value is None or value == "":
        return 0
    try:
        f = float(str(value).replace(",", "").strip())
    except (ValueError, TypeError):
        return 0
    return int(f) if f == int(f) else f


def to_str(value):
    if value is None:
        return ""
    return str(value).strip()


def to_float(value):
    """None/空 -> 0.0"""
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (ValueError, TypeError):
        return 0.0


def parse_date(value, context=""):
    """
    需求日期归一化 -> "YYYY-MM-DD"。
    遇到 "2026-09-30~2026-10-01" 这种区间写法，取开始日期
    （需求开始到岗的日期；基本信息透视表对区间行的记法本身就不一致，
    有记在开始日的、有两天都记的，这里统一按开始日期归属，并在
    warning 里记一笔，透视表逐日比对时排除区间行）。
    解析失败返回 "" 并记 warning（调用方决定是否丢弃该行）。
    """
    if value is None or str(value).strip() == "":
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    m = re.match(r"^(\d{4}-\d{2}-\d{2})\s*~\s*(\d{4}-\d{2}-\d{2})$", s)
    if m:
        warn(f"日期区间写法已取开始日期: {s} -> {m.group(1)} ({context})")
        return m.group(1)
    m = re.match(r"^(\d{4}-\d{2}-\d{2})$", s)
    if m:
        return m.group(1)
    # Excel 数字日期
    try:
        f = float(s)
        from datetime import datetime, timedelta
        base = datetime(1899, 12, 30)
        return (base + timedelta(days=int(f))).strftime("%Y-%m-%d")
    except (ValueError, TypeError):
        pass
    warn(f"无法解析的日期: {s} ({context})")
    return ""


# ------------------------------------------------------------------
# 工种归一化
# ------------------------------------------------------------------
# 需求详情里是 "Cherry Picker I"（带 Level），发单详情里是 "Cherry Picker"。
# 跨表对比时用归一化名，原始名保留在「原始工种」列里备查。
# 只去掉末尾的罗马数字 Level（I/II/III），不合并不同工种
#（比如不会把「卸柜员」和 Lumper 合并）。

_LEVEL_SUFFIX = re.compile(r"\s+[IVX]{1,3}$")


def norm_job(name):
    original = to_str(name)
    normalized = _LEVEL_SUFFIX.sub("", original).strip()
    return normalized or original, original


# ------------------------------------------------------------------
# 文件分类
# ------------------------------------------------------------------

def classify_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    sheets = set(wb.sheetnames)
    kind = "unknown"
    if "需求详情" in sheets and "发单详情" in sheets:
        kind = "demand_pool"
    elif "派遣人员名单" in sheets:
        kind = "dispatch_list"
    elif "考勤" in sheets:
        kind = "attendance"
    elif "临时工派遣" in sheets:
        kind = "temp_dispatch"
    elif "Sheet1" in sheets:
        # 人员到岗表：Sheet1，有 NO SHOW 和 Not Fufilled Reason 列
        ws = wb["Sheet1"]
        hdr = " ".join(
            str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)
        ).upper()
        if "NO SHOW" in hdr and "FUFILLED" in hdr:
            kind = "arrival_tracking"
    wb.close()
    return kind


# ------------------------------------------------------------------
# 需求池文件：需求详情 / 发单详情 / 基本信息
# ------------------------------------------------------------------

def read_demand_detail(ws):
    """需求详情 -> list[dict]，一行一条需求（日期+仓+组+工种+班次）"""
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if r[0] is None and r[1] is None:
            continue
        date_raw = r[3]
        is_range = isinstance(date_raw, str) and "~" in date_raw
        date = parse_date(date_raw, context=f"需求详情第{i+1}行")
        if not date:
            warn(f"需求详情第{i+1}行日期为空，已跳过")
            continue
        job_norm, job_orig = norm_job(r[4])
        rows.append({
            "_is_range": is_range,  # 校验用：透视表逐日比对时排除
            "需求区域": to_str(r[0]),
            "需求仓": to_str(r[1]),
            "需求组": to_str(r[2]),
            "需求日期": date,
            "工种": job_norm,
            "原始工种": job_orig,
            "班次": to_str(r[5]),
            "需求人数": to_int(r[6]),
            "待发单人数": to_int(r[7]),
            "已发单人数": to_int(r[8]),
            "待派遣人数": to_int(r[9]),
            "已派遣人数": to_int(r[10]),
            "用工要求": to_str(r[11]),
            "建议派遣人员": to_str(r[12]),
            "关联申请编号": to_str(r[13]),
        })
    return rows


def read_alloc_detail(ws):
    """发单详情 -> list[dict]，一行一条供应商分配"""
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if r[0] is None:
            continue
        date = parse_date(r[6], context=f"发单详情第{i+1}行")
        if not date:
            warn(f"发单详情第{i+1}行日期为空，已跳过")
            continue
        job_norm, job_orig = norm_job(r[7])
        rows.append({
            "供应商": to_str(r[0]),
            "供应商名称": to_str(r[0]),
            "供应商ID": to_str(r[1]),
            "类型": to_str(r[2]),
            "需求区域": to_str(r[3]),
            "需求仓": to_str(r[4]),
            "需求组": to_str(r[5]),
            "需求日期": date,
            "工种": job_norm,
            "原始工种": job_orig,
            "班次": to_str(r[8]),
            # 口径：已发单人数 = 分给该供应商的人数（供应商的"需求"）
            "供应商需派遣人数": to_int(r[9]),
            "供应商未派遣人数": to_int(r[10]),
            "供应商已派遣人数": to_int(r[11]),
            "用工要求": to_str(r[12]),
            "建议派遣人员": to_str(r[13]),
        })
    return rows


def read_basic_info(ws):
    """
    基本信息（透视表）-> {(区域,仓,组,日期): {需求总人数, 补员, ...}}
    只用于交叉校验，不进 dashboard。
    """
    # 第0行是日期头（每7列一个日期），第1行是子列头
    header0, header1 = None, None
    info = {}
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            header0 = [to_str(v) for v in r]
            continue
        if i == 1:
            header1 = [to_str(v) for v in r]
            continue
        if r[0] is None:
            continue
        area, wh, grp = to_str(r[0]), to_str(r[1]), to_str(r[2])
        # 每7列一组：需求总人数|补员|待发单人数|已发单人数|待派遣人数|已派遣人数|需求完成率
        for start in range(3, len(header0) - 1, 7):
            date_label = header0[start]  # 如 "09-01(周二)"
            m = re.match(r"(\d{2})-(\d{2})", date_label)
            if not m:
                continue
            date = f"2026-{m.group(1)}-{m.group(2)}"
            vals = [to_int(r[start + k]) if start + k < len(r) else 0
                    for k in range(6)]
            if any(vals):
                info[(area, wh, grp, date)] = {
                    "需求总人数": vals[0],
                    "补员": vals[1],
                    "待发单人数": vals[2],
                    "已发单人数": vals[3],
                    "待派遣人数": vals[4],
                    "已派遣人数": vals[5],
                }
    return info


# ------------------------------------------------------------------
# 派遣名单文件
# ------------------------------------------------------------------

def read_dispatch_list(ws):
    """
    派遣人员名单 -> list[dict]。
    前12列是需求信息（表头在第0行），后5列是人员信息
    （表头在第1行：供应商|供应商ID|员工姓名|员工工号|派遣开始日期）。
    """
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i <= 1:
            continue
        if r[0] is None and r[1] is None:
            continue
        date = parse_date(r[4], context=f"派遣名单第{i+1}行")
        job_norm, job_orig = norm_job(r[5])
        rows.append({
            "序号": to_str(r[0]),
            "需求区域": to_str(r[1]),
            "需求仓": to_str(r[2]),
            "需求组": to_str(r[3]),
            "需求日期": date,
            "工种": job_norm,
            "原始工种": job_orig,
            "班次": to_str(r[6]),
            "用工要求": to_str(r[7]),
            "建议派遣人员": to_str(r[8]),
            "供应商需派遣人数": to_int(r[9]),
            "供应商未派遣人数": to_int(r[10]),
            "供应商已派遣人数": to_int(r[11]),
            "供应商": to_str(r[12]),
            "供应商ID": to_str(r[13]),
            "员工姓名": to_str(r[14]),
            "员工工号": to_str(r[15]),
            "派遣开始日期": parse_date(r[16], context=f"派遣名单第{i+1}行派遣开始日期")
                                  if r[16] else "",
        })
    return rows


# ------------------------------------------------------------------
# 考勤表（22 列格式）
# ------------------------------------------------------------------
# A~V：供应商名称 供应商ID 姓名 工号 工种 员工类型 区域 仓库 组 考勤类型
#      考勤日期 记录生成时间 白班工作时长 下午班工作时长 晚班工作时长
#      工作时长 加班时长 时长总计 班次 考勤方案 首打卡时间 末打卡时间
#
# 一行 = 一人一天一条考勤记录。表头对不上（旧的空导出只有一列）
# 时返回 None，调用方按空源处理。


def read_attendance(ws):
    header = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            header = [to_str(v) for v in row]
            break
    if not header or "考勤日期" not in header:
        return None
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        d = dict(zip(header, row))
        att_date = parse_date(d.get("考勤日期"), context=f"考勤第{i+1}行")
        if not att_date:
            warn(f"考勤第{i+1}行日期为空，已跳过")
            continue
        job_norm, _ = norm_job(d.get("工种"))
        rows.append({
            "供应商": to_str(d.get("供应商名称")) or "未分类",
            "供应商ID": to_str(d.get("供应商ID")),
            "姓名": to_str(d.get("姓名")),
            "工号": to_str(d.get("工号")),
            "工种": job_norm or "未分类",
            "仓库": to_str(d.get("仓库")) or "未分类",
            "组": to_str(d.get("组")),
            "考勤类型": to_str(d.get("考勤类型")),
            "考勤日期": att_date,
            "工作时长": round(to_float(d.get("工作时长")), 2),
            "加班时长": round(to_float(d.get("加班时长")), 2),
            "首打卡": to_str(d.get("首打卡时间")),
            "末打卡": to_str(d.get("末打卡时间")),
        })
    return rows


# ------------------------------------------------------------------
# 临时工派遣表（24 列）
# ------------------------------------------------------------------
# 派遣编码 姓名 工号 编号 员工状态 服务商名称 服务商编码 供应商名称
# 供应商ID 区域 物理仓 组 派遣开始日期 派遣结束日期 报价方案名称
# 报价方案编码 工种 班次 考勤方案 状态 创建人 创建时间 更新人 更新时间
#
# 一行 = 一人一天一条派遣记录。状态 ∈ {已确认, 试工不通过, 待确认}。
# 表头对不上时返回 None，调用方按空源处理。


def read_temp_dispatch(ws):
    header = [to_str(v) for v in next(ws.iter_rows(values_only=True))]
    if "派遣编码" not in header or "状态" not in header:
        return None
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        d = dict(zip(header, row))
        if d.get("派遣编码") is None and not to_str(d.get("姓名")):
            continue
        job_norm, _ = norm_job(d.get("工种"))
        rows.append({
            "派遣编码": to_str(d.get("派遣编码")),
            "姓名": to_str(d.get("姓名")),
            "工号": to_str(d.get("工号")),
            "员工状态": to_str(d.get("员工状态")),
            "供应商": to_str(d.get("供应商名称")) or "未分类",
            "供应商ID": to_str(d.get("供应商ID")),
            "仓库": to_str(d.get("物理仓")) or "未分类",
            "组": to_str(d.get("组")),
            "派遣开始日期": parse_date(d.get("派遣开始日期"),
                                      context=f"临时工派遣第{i+1}行"),
            "派遣结束日期": parse_date(d.get("派遣结束日期"),
                                      context=f"临时工派遣第{i+1}行"),
            "工种": job_norm or "未分类",
            "班次": to_str(d.get("班次")),
            "状态": to_str(d.get("状态")),
        })
    return rows


def read_arrival_tracking(ws):
    """
    人员到岗表 (Sheet1) -> list[dict]，一行一条拟到岗记录。

    用户口径：
      - NO SHOW 未勾选（0/空）= 到岗
      - 到岗 且 不符合要求原因 为空 = 已接受
      - NO SHOW 勾选（1）= no show
      - 不符合要求原因 非空 = not qualify（被退回）
    """
    header = [to_str(v) for v in next(ws.iter_rows(values_only=True))]
    hdr_up = [h.upper() for h in header]
    if not any("NO SHOW" in h for h in hdr_up):
        return None
    # 列定位（表头是中英双语，取包含关键词的列）
    def col(*keys):
        for i, h in enumerate(hdr_up):
            if all(k in h for k in keys):
                return i
        return None
    i_loc = col("LOCATION")
    i_date = col("REPORT DATE")
    i_pos = col("POSITION")
    i_agency = col("AGENCY")
    i_name = col("EMPLOYEE NAME")
    i_id = col("OTWS ID")
    i_noshow = next((i for i, h in enumerate(hdr_up) if "NO SHOW" in h), None)
    i_reason = next((i for i, h in enumerate(hdr_up) if "FUFILLED" in h), None)
    rows = []
    for n, row in enumerate(ws.iter_rows(values_only=True)):
        if n == 0:
            continue
        name = to_str(row[i_name]) if i_name is not None else ""
        wid = to_str(row[i_id]) if i_id is not None else ""
        # 用户口径：有 OTWS ID 才算需求日到场
        if not wid:
            continue
        noshow_raw = to_str(row[i_noshow]).strip() if i_noshow is not None else ""
        reason = to_str(row[i_reason]).strip() if i_reason is not None else ""
        is_noshow = noshow_raw == "1"
        arrived = not is_noshow
        accepted = arrived and not reason
        rows.append({
            "仓库": to_str(row[i_loc]).strip() if i_loc is not None else "未分类",
            "到岗日期": parse_date(row[i_date], context=f"人员到岗第{n+1}行")
                        if i_date is not None else None,
            "职位": to_str(row[i_pos]) if i_pos is not None else "",
            "机构": to_str(row[i_agency]).strip() if i_agency is not None else "未分类",
            "姓名": name,
            "工号": wid,
            "是否NoShow": is_noshow,
            "不符合要求原因": reason,
            "到岗": arrived,
            "已接受": accepted,
        })
    return rows


# ------------------------------------------------------------------
# 入口：读整个 raw 目录
# ------------------------------------------------------------------

def load_raw(raw_dir):
    """
    返回 {
      "demand": [...],        # 需求详情行
      "alloc": [...],         # 发单详情行
      "dispatch": [...],      # 派遣名单行（含未派遣的空行）
      "basic_info": {...},    # 基本信息透视（仅校验用）
      "files": {path: kind},  # 识别到的文件
      "empty_sources": [...], # 有表头但零数据的文件
    }
    """
    reset_warnings()
    raw_dir = Path(raw_dir)
    out = {
        "demand": [], "alloc": [], "dispatch": [], "attendance": [],
        "temp_dispatch": [], "arrival_tracking": [],
        "basic_info": {}, "files": {}, "empty_sources": [],
    }
    for path in sorted(raw_dir.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        kind = classify_workbook(path)
        out["files"][str(path)] = kind
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
        if kind == "demand_pool":
            out["demand"].extend(read_demand_detail(wb["需求详情"]))
            out["alloc"].extend(read_alloc_detail(wb["发单详情"]))
            for k, v in read_basic_info(wb["基本信息"]).items():
                out["basic_info"][k] = v
        elif kind == "dispatch_list":
            out["dispatch"].extend(read_dispatch_list(wb["派遣人员名单"]))
        elif kind == "attendance":
            att = (read_attendance(wb["考勤"])
                   if "考勤" in wb.sheetnames else None)
            if att:
                out["attendance"].extend(att)
            else:
                # 有表头、零数据行 -> 记下来，后面产出空数据集
                out["empty_sources"].append((str(path), kind))
                warn(f"{path.name} 只有表头、没有数据行，考勤/OT相关指标将为空")
        elif kind == "temp_dispatch":
            td = (read_temp_dispatch(wb["临时工派遣"])
                  if "临时工派遣" in wb.sheetnames else None)
            if td:
                out["temp_dispatch"].extend(td)
            else:
                # 有表头、零数据行 -> 记下来，后面产出空数据集
                out["empty_sources"].append((str(path), kind))
                warn(f"{path.name} 只有表头、没有数据行，临时工派遣相关指标将为空")
        elif kind == "arrival_tracking":
            at = (read_arrival_tracking(wb["Sheet1"])
                  if "Sheet1" in wb.sheetnames else None)
            if at:
                out["arrival_tracking"].extend(at)
            else:
                out["empty_sources"].append((str(path), kind))
                warn(f"{path.name} 只有表头、没有数据行，到岗/接受指标将为空")
        else:
            warn(f"无法识别的文件已跳过: {path.name}")
        wb.close()
    return out
